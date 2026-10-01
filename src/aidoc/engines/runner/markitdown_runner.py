"""MarkItDown runner. STDLIB + markitdown (+ python-pptx) only; never import aidoc."""
from __future__ import annotations

import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
import _proto

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}
IMG_LINK = re.compile(r"!\[([^\]]*)\]\(([^)\s]+)\)")
AUDIO_EXTS = {".mp3", ".wav", ".m4a"}


def docx_media_in_order(path):
    """(filename, bytes) of every embedded picture, in document order."""
    with zipfile.ZipFile(path) as z:
        rels = ET.fromstring(z.read("word/_rels/document.xml.rels"))
        rid2target = {r.get("Id"): r.get("Target") for r in rels.findall("rel:Relationship", NS)}
        doc = ET.fromstring(z.read("word/document.xml"))
        out = []
        for blip in doc.iter(f"{{{NS['a']}}}blip"):
            rid = blip.get(f"{{{NS['r']}}}embed")
            target = rid2target.get(rid)
            if target:
                name = "word/" + target if not target.startswith("/") else target[1:]
                try:
                    out.append((Path(target).name, z.read(name)))
                except KeyError:
                    continue
        return out


def replace_nth_images(md, replacements):
    it = iter(replacements)

    def sub(m):
        try:
            return f"![{m.group(1)}]({next(it)})"
        except StopIteration:
            return m.group(0)
    return IMG_LINK.sub(sub, md)


def slide_markers_to_pages(md):
    return re.sub(r"<!-- Slide number: (\d+) -->", r"<!-- page: \1 -->", md)


def pptx_pictures_in_order(path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    out = []

    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(sh.shapes)
            elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                out.append((sh.image.ext, sh.image.blob))
    for slide in Presentation(path).slides:
        walk(slide.shapes)
    return out


def md_normalize(text):
    """MarkItDown's own post-processing of every conversion (rstrip lines, collapse 3+ newlines)."""
    text = "\n".join(line.rstrip() for line in re.split(r"\r?\n", text))
    return re.sub(r"\n{3,}", "\n\n", text)


def split_formfeed(text, n_pages):
    """pdfminer puts a form feed after every page: split, drop one trailing empty part; None if the count differs."""
    parts = text.split("\f")
    if parts and not parts[-1].strip():
        parts = parts[:-1]
    return parts if len(parts) == n_pages else None


def join_like_markitdown(pages, form_path):
    """MarkItDown's own join: form path = non-empty chunks joined by blank lines; pdfminer path = original text."""
    if form_path:
        return "\n\n".join(p for p in pages if p).strip()
    return "\f".join(pages) + "\f"


def pdf_pages(path):
    """Spec 2026-10-01 §6.3: MarkItDown's PDF conversion, page by page, with its own primitives. None when the
    pages, joined back MarkItDown's way, do not reproduce `MarkItDown().convert()` exactly (self-check)."""
    import io

    import pdfminer.high_level
    import pdfplumber
    from markitdown import MarkItDown
    from markitdown.converters import _pdf_converter as P

    data = io.BytesIO(Path(path).read_bytes())
    chunks, form = [], 0
    with pdfplumber.open(data) as pdf:
        for page in pdf.pages:
            content = P._extract_form_content_from_words(page)
            if content is not None:
                form += 1
                chunks.append(content if content.strip() else "")
            else:
                chunks.append((page.extract_text() or "").strip())
            page.close()
    if form == 0:
        data.seek(0)
        text = pdfminer.high_level.extract_text(data)
        pages = split_formfeed(text, len(chunks))
        if pages is None:
            return None
        joined = text
    else:
        pages = chunks
        joined = join_like_markitdown(pages, form_path=True)
    ref = MarkItDown(enable_plugins=False).convert(str(path)).markdown
    if md_normalize(P._merge_partial_numbering_lines(joined)) != ref:
        return None
    return pages


def handle(req):
    from markitdown import MarkItDown
    src = Path(req["src"])
    out = Path(req["out_dir"])
    (out / "images").mkdir(parents=True, exist_ok=True)
    ext = src.suffix.lower()
    if ext in AUDIO_EXTS and not req["engine_opts"].get("allow_online_audio"):
        raise _proto.RunnerError("input", "audio_disabled")
    try:
        md = MarkItDown(enable_plugins=False).convert(str(src)).markdown
    except Exception as e:  # unsupported / unreadable input
        name = type(e).__name__
        if name in ("UnsupportedFormatException", "FileConversionException"):
            raise _proto.RunnerError("input", f"{name}: {e}")
        raise
    images, has_pages, method = [], False, None
    if ext == ".pdf" or req.get("kind") == "pdf":
        try:
            pages = pdf_pages(src)
        except Exception as e:  # noqa: BLE001  internals changed: no page map (-> page_map_incomplete -> low)
            _proto.log(f"markitdown per-page extraction failed: {type(e).__name__}: {e}")
            pages = None
        if pages is not None:
            md = "".join(f"<!-- page: {i} -->\n\n" + (md_normalize(p).strip("\n") + "\n\n" if p.strip() else "")
                         for i, p in enumerate(pages, 1))
            has_pages, method = True, "markitdown_per_page"
        else:
            _proto.log("markitdown per-page self-check failed: output has no page markers")
            method = "none"
    if ext == ".docx":
        media = docx_media_in_order(src)
    elif ext == ".pptx":
        media = [(f"pic{i+1}.{e}", b) for i, (e, b) in enumerate(pptx_pictures_in_order(src))]
        md = slide_markers_to_pages(md)
        has_pages = "<!-- page: " in md
    else:
        media = []
    names = []
    for i, (name, blob) in enumerate(media):
        p = out / "images" / f"img_{i+1}{Path(name).suffix.lower() or '.png'}"
        p.write_bytes(blob)
        images.append(str(p))
        names.append(f"images/{p.name}")
    md = replace_nth_images(md, names)
    md_path = out / "out.md"
    md_path.write_text(md, encoding="utf-8")
    _proto.progress(1, 1)
    return {"markdown_path": str(md_path), "images": images, "has_page_markers": has_pages,
            "page_count": None, "first_table": None, "last_table": None, "page_map_method": method}


if __name__ == "__main__":
    import markitdown  # noqa: F401  (load the library before READY so the timeout clock excludes imports)
    _proto.serve(handle)
