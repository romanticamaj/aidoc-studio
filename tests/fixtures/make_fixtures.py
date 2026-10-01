"""Regenerate all test fixtures. Run: uv run python tests/fixtures/make_fixtures.py"""
from __future__ import annotations

import random
import re
import shutil
from pathlib import Path

import fitz  # PyMuPDF

HERE = Path(__file__).parent
SAMPLES = HERE.parent.parent / "src" / "aidoc" / "engines" / "runner" / "samples"
CJK_FONTFILE = None
for cand in [r"C:\Windows\Fonts\msjh.ttc", r"C:\Windows\Fonts\msjhbd.ttc",
             "/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"]:
    if Path(cand).exists():
        CJK_FONTFILE = cand
        break

ZH = "這是一份用來測試文件轉換流程的繁體中文段落。內容包含標題、表格與圖片，用於驗證引擎的中文辨識能力。"
EN = ("This is an English paragraph used to test the document conversion pipeline. "
      "It contains enough words to exceed the per-page threshold.")

_FONT = None


def _font(page=None):
    global _FONT
    if _FONT is None:
        _FONT = fitz.Font(fontfile=CJK_FONTFILE) if CJK_FONTFILE else fitz.Font("cjk")
    return _FONT


def _wrap(text, font, size, max_w):
    """Greedy wrap by measured width; Latin words are kept whole, CJK breaks anywhere."""
    tokens = re.findall(r"[A-Za-z0-9.,;:'()/-]+ ?| |.", text)
    lines, cur = [], ""
    for tok in tokens:
        if cur and font.text_length(cur + tok.rstrip(), fontsize=size) > max_w:
            lines.append(cur.rstrip())
            cur = tok.lstrip()
        else:
            cur += tok
    if cur.strip():
        lines.append(cur.rstrip())
    return lines


def _write_paragraphs(page, texts, y=72, size=11):
    tw = fitz.TextWriter(page.rect)
    font = _font(page)
    max_w = page.rect.width - 144
    for t in texts:
        for line in _wrap(t, font, size, max_w):
            tw.append((72, y), line, font=font, fontsize=size)
            y += size + 5
        y += 8
    tw.write_text(page)
    return y


def _draw_table(page, y, rows=4, cols=3, w=450, rh=22, label="R"):
    x0 = 72
    sh = page.new_shape()
    for r in range(rows + 1):
        sh.draw_line((x0, y + r * rh), (x0 + w, y + r * rh))
    for c in range(cols + 1):
        sh.draw_line((x0 + c * w / cols, y), (x0 + c * w / cols, y + rows * rh))
    sh.finish(width=0.8, color=(0, 0, 0))
    sh.commit()
    tw = fitz.TextWriter(page.rect)
    font = _font(page)
    for r in range(rows):
        for c in range(cols):
            tw.append((x0 + c * w / cols + 4, y + r * rh + 15), f"{label}{r}C{c} 值", font=font, fontsize=10)
    tw.write_text(page)


def _save(doc, path):
    try:
        doc.subset_fonts()          # keep committed fixtures small (msjh.ttc is ~20 MB)
    except Exception:  # noqa: BLE001, S110  subsetting is an optimisation only
        pass
    doc.save(path, garbage=4, deflate=True)
    doc.close()


def make_text_pdf(path):
    doc = fitz.open()
    for n in range(3):
        page = doc.new_page()
        y = _write_paragraphs(page, [f"Page {n+1} 標題", ZH * 3, EN * 2])
        if n == 1:
            _draw_table(page, y + 10)
    _save(doc, path)


def _render_image_page(texts, dpi=150):
    doc = fitz.open()
    page = doc.new_page()
    _write_paragraphs(page, texts, size=14)
    pix = page.get_pixmap(dpi=dpi)
    png = pix.tobytes("png")
    doc.close()
    return png


def make_scanned_pdf(path, pages_texts):
    doc = fitz.open()
    for texts in pages_texts:
        png = _render_image_page(texts)
        page = doc.new_page()
        page.insert_image(page.rect, stream=png)
    _save(doc, path)


def _typeset(page, x, y, parts, size=18):
    """Draw a display equation from (text, dy, scale) runs: baseline shifts for sub/superscripts."""
    tw = fitz.TextWriter(page.rect)
    font = _font(page)
    for text, dy, scale in parts:
        fs = size * scale
        tw.append((x, y + dy * size), text, font=font, fontsize=fs)
        x += font.text_length(text, fontsize=fs) + 1
    tw.write_text(page)
    return x


def _fraction(page, x, y, num, den, size=18):
    font = _font(page)
    w = max(font.text_length(num, fontsize=size), font.text_length(den, fontsize=size)) + 6
    tw = fitz.TextWriter(page.rect)
    tw.append((x + (w - font.text_length(num, fontsize=size)) / 2, y - 0.55 * size), num, font=font, fontsize=size)
    tw.append((x + (w - font.text_length(den, fontsize=size)) / 2, y + 0.75 * size), den, font=font, fontsize=size)
    tw.write_text(page)
    page.draw_line((x, y - 0.3 * size), (x + w, y - 0.3 * size), width=0.9)
    return x + w + 2


def _display_equations(page, y):
    """Typeset display equations (integral with limits, fraction, summation, superscripts)."""
    x = 150
    x = _typeset(page, x, y, [("∫", 0.15, 1.6), ("1", -0.75, 0.6)])
    _typeset(page, x - 14, y + 12, [("0", 0, 0.6)])
    x = _typeset(page, x, y, [(" x", 0, 1), ("2", -0.45, 0.6), (" dx = ", 0, 1)])
    _fraction(page, x, y, "1", "3")
    y += 70
    x = 150
    x = _typeset(page, x, y, [("f(x) = ", 0, 1), ("Σ", 0.1, 1.5)])
    _typeset(page, x - 22, y - 26, [("∞", 0, 0.6)])
    _typeset(page, x - 26, y + 14, [("n=0", 0, 0.6)])
    _typeset(page, x, y, [(" a", 0, 1), ("n", 0.3, 0.6), (" x", 0, 1), ("n", -0.45, 0.6)])
    y += 60
    x = _typeset(page, 150, y, [("E = mc", 0, 1), ("2", -0.45, 0.6), (",   a", 0, 1), ("2", -0.45, 0.6),
                               (" + b", 0, 1), ("2", -0.45, 0.6), (" = c", 0, 1), ("2", -0.45, 0.6)])
    y += 60
    x = _typeset(page, 150, y, [("x = ", 0, 1)])
    _fraction(page, x, y, "−b ± √(b² − 4ac)", "2a")


def make_formula_pdf(path):
    doc = fitz.open()
    for n in range(2):
        page = doc.new_page()
        _write_paragraphs(page, [EN, ZH])
        _display_equations(page, 230)
        page.insert_text((72, 560), "abgdSyxw " * 8, fontname="symb", fontsize=12)
        tw = fitz.TextWriter(page.rect)
        tw.append((72, 600), "∑ ∫ ∂ √ ∞ ≈ ≠ ≤ ≥ ± × ÷ ∈ ∀ ∃ α β γ " * 2, font=_font(page), fontsize=12)
        tw.write_text(page)
    _save(doc, path)


def make_twocol_pdf(path):
    doc = fitz.open()
    for n in range(2):
        page = doc.new_page()
        tw = fitz.TextWriter(page.rect)
        font = _font(page)
        for col_x in (60, 320):
            y = 72
            for _ in range(28):
                tw.append((col_x, y), (EN + ZH)[:38], font=font, fontsize=10)
                y += 16
        tw.write_text(page)
    _save(doc, path)


def make_big_pdf(path, n=45):
    doc = fitz.open()
    for i in range(n):
        page = doc.new_page()
        _write_paragraphs(page, [f"第 {i+1} 頁 / Page {i+1}", ZH * 2, EN])
    _save(doc, path)


def make_span_margin_pdf(path, n=45, split_after=40):
    """P2 verifier I2: a table that runs from the bottom margin of page 40 onto the top margin of page 41 (the
    40-page segment boundary), with ordinary 1-inch margins and a page number in the footer."""
    doc = fitz.open()
    for i in range(1, n + 1):
        page = doc.new_page()                                  # A4-ish default 595 x 842
        bottom = page.rect.height - 72
        if i == split_after:
            y = _write_paragraphs(page, [f"Section {i} 表格開始", ZH, EN])
            rows = int((bottom - (y + 10)) // 22)
            _draw_table(page, bottom - rows * 22, rows=rows, label="A")
        elif i == split_after + 1:
            _draw_table(page, 72, rows=5, label="B")
            _write_paragraphs(page, [ZH, EN], y=72 + 5 * 22 + 24)
        else:
            _write_paragraphs(page, [f"Section {i}", ZH * 2, EN])
        tw = fitz.TextWriter(page.rect)
        tw.append((page.rect.width / 2 - 10, page.rect.height - 36), f"- {i} -", font=_font(page), fontsize=9)
        tw.write_text(page)
    _save(doc, path)


def make_blank_pdf(path):
    doc = fitz.open()
    doc.new_page()
    doc.save(path)
    doc.close()


def make_encrypted_pdf(path):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "secret")
    doc.save(path, encryption=fitz.PDF_ENCRYPT_AES_256, user_pw="pw", owner_pw="pw")
    doc.close()


def make_corrupt_pdf(path):
    random.seed(1)
    Path(path).write_bytes(b"%PDF-1.7\n" + bytes(random.getrandbits(8) for _ in range(2048)))


def make_png(path):
    Path(path).write_bytes(_render_image_page([ZH, EN]))


def make_docx(path, png):
    import docx
    import docx.shared
    d = docx.Document()
    d.add_heading("測試文件 Sample", 1)
    d.add_paragraph(ZH)
    d.add_heading("Section 2", 2)
    d.add_paragraph(EN)
    t = d.add_table(rows=3, cols=3)
    for r in range(3):
        for c in range(3):
            t.cell(r, c).text = f"R{r}C{c}"
    d.add_picture(str(png), width=docx.shared.Inches(2))
    d.save(path)


def make_pptx(path, png):
    from pptx import Presentation
    from pptx.util import Inches
    p = Presentation()
    for i in range(3):
        s = p.slides.add_slide(p.slide_layouts[1])
        s.shapes.title.text = f"投影片 {i+1} Slide {i+1}"
        s.placeholders[1].text = ZH + " " + EN
        if i == 1:
            s.shapes.add_picture(str(png), Inches(1), Inches(3), width=Inches(3))
    p.save(path)


def make_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["名稱", "數量", "Price"])
    for i in range(1, 6):
        ws.append([f"item{i}", i, i * 1.5])
    ws2 = wb.create_sheet("第二頁")
    ws2.append(["a", "b"])
    ws2.append([1, 2])
    wb.save(path)


def make_html(path):
    Path(path).write_text(
        f"<html><body><h1>標題 Title</h1><p>{ZH}</p><p>{EN}</p><table><tr><th>a</th><th>b</th></tr>"
        f"<tr><td>1</td><td>2</td></tr></table></body></html>", encoding="utf-8")


_FURNITURE_WORDS = ("patient joint muscle tendon ligament examination movement passive active resisted painful "
                    "capsular pattern lesion treatment injection friction massage manipulation traction posture "
                    "shoulder elbow wrist hip knee ankle spine nerve root dura sign test history onset "
                    "gradual sudden morning evening stiffness swelling warmth weakness range limited full").split()


def _furniture_body(page_no: int, n_words: int = 130) -> str:
    rng = random.Random(1000 + page_no)
    words = [rng.choice(_FURNITURE_WORDS) for _ in range(n_words)]
    for i in range(0, n_words, 12):                    # page-specific tokens keep every page's text unique
        words[i] = f"{words[i]}{page_no}{chr(97 + (i // 12) % 26)}"
    return f"Body text of page {page_no}: " + " ".join(words) + "."


def _helv_lines(page, text, x, y, size=10, max_w=450):
    font = fitz.Font("helv")
    for line in _wrap(text, font, size, max_w):
        page.insert_text((x, y), line, fontname="helv", fontsize=size)
        y += size + 4
    return y


def make_paged_furniture_pdf(path):
    """8 pages (A4): running header + footer page number on every non-blank page; page 2 is a table of contents
    with dot leaders; page 4 is blank; a paragraph starts at the bottom of page 5 and continues on page 6; page 7
    has a drawn figure + caption. Body pages carry >= 600 chars of unique English text."""
    doc = fitz.open()
    for n in range(1, 9):
        page = doc.new_page(width=595, height=842)
        if n == 4:
            continue                                   # blank: no text, no drawing
        page.insert_text((72, 40), "Chapter 2 - Testing Pages", fontname="helv", fontsize=9)
        page.insert_text((290, 815), str(n), fontname="helv", fontsize=9)
        if n == 2:
            page.insert_text((72, 90), "Contents", fontname="helv", fontsize=16)
            y = 120
            for title, pg in (("Introduction", 3), ("Examination of the joint", 5), ("Spanning paragraph", 6),
                              ("Figures", 7), ("Summary", 8)):
                page.insert_text((72, y), f"{title} {'.' * (60 - len(title))} {pg}", fontname="helv", fontsize=11)
                y += 20
            _helv_lines(page, _furniture_body(n, 100), 72, y + 20)
            continue
        y = _helv_lines(page, _furniture_body(n), 72, 80)
        if n == 5:
            # a paragraph that starts at the bottom of page 5 and continues on page 6
            _helv_lines(page, "Spanning paragraph begins here and keeps describing the capsular pattern of the "
                              "shoulder joint while the page runs out of room at the very bottom margin of the",
                        72, 770)
        if n == 6:
            _helv_lines(page, "sheet, so the reader continues on the next page where the spanning paragraph "
                              "finally ends with a full stop.", 72, 60)
            _helv_lines(page, _furniture_body(60 + n, 40), 72, y + 20)
        if n == 7:
            sh = page.new_shape()
            sh.draw_rect(fitz.Rect(150, y + 20, 450, y + 220))
            sh.finish(color=(0, 0, 0), fill=(0.3, 0.5, 0.8))
            sh.commit()
            page.insert_text((150, y + 240), "Figure 7.1 A drawn box", fontname="helv", fontsize=10)
    _save(doc, path)


BROKEN_SENTENCE = ("Wouldn't it be dreamy if there was a book on Android",
                   "development that could turn me into an expert while",
                   "keeping me engaged and entertained?")


def make_broken_tounicode_pdf(path):
    """2 pages; page 1 has three lines in a Type0/Identity-H font whose /ToUnicode was deleted (T-001), page 2 uses
    that font only for a footer. The text layer of those lines is glyph ids."""
    doc = fitz.open()
    for n in (1, 2):
        page = doc.new_page(width=595, height=842)
        page.insert_font(fontname="K", fontbuffer=fitz.Font("cjk").buffer)
        page.insert_text((72, 72), f"Broken ToUnicode fixture - page {n}", fontname="helv", fontsize=16)
        y = 110
        for i in range(6):
            page.insert_text((72, y), f"Normal Helvetica line {i + 1} on page {n}: the quick brown fox jumps over "
                                      f"the lazy dog.", fontname="helv", fontsize=11)
            y += 18
        if n == 1:
            y += 20
            for line in BROKEN_SENTENCE:
                page.insert_text((72, y), line, fontname="K", fontsize=14)
                y += 22
        else:
            page.insert_text((290, 815), "7", fontname="K", fontsize=9)
    try:
        doc.subset_fonts()                             # keep the committed file small (the cjk font is ~4 MB)
    except Exception:  # noqa: BLE001, S110  subsetting is an optimisation only
        pass
    for xref in range(1, doc.xref_length()):
        try:
            if doc.xref_get_key(xref, "Subtype")[1] == "/Type0":
                doc.xref_set_key(xref, "ToUnicode", "null")
        except Exception:  # noqa: BLE001, S112  non-dict objects
            continue
    doc.save(path, garbage=4, deflate=True)
    doc.close()


def main():
    HERE.mkdir(exist_ok=True)
    SAMPLES.mkdir(parents=True, exist_ok=True)
    make_text_pdf(HERE / "text.pdf")
    make_scanned_pdf(HERE / "scanned_cht.pdf", [[ZH * 4], [ZH * 4]])
    make_scanned_pdf(HERE / "scanned_mixed.pdf", [[ZH, EN, ZH], [EN, ZH, EN]])
    make_scanned_pdf(HERE / "scanned_en.pdf", [[EN * 4]])
    make_formula_pdf(HERE / "formula.pdf")
    make_twocol_pdf(HERE / "twocol.pdf")
    make_big_pdf(HERE / "big.pdf")
    make_blank_pdf(HERE / "blank.pdf")
    make_span_margin_pdf(HERE / "span_margin.pdf")
    make_paged_furniture_pdf(HERE / "paged_furniture.pdf")
    make_broken_tounicode_pdf(HERE / "broken_tounicode.pdf")
    make_encrypted_pdf(HERE / "encrypted.pdf")
    make_corrupt_pdf(HERE / "corrupt.pdf")
    make_png(HERE / "page.png")
    make_docx(HERE / "sample.docx", HERE / "page.png")
    make_pptx(HERE / "sample.pptx", HERE / "page.png")
    make_xlsx(HERE / "sample.xlsx")
    make_html(HERE / "sample.html")
    (HERE / "bad.exe").write_bytes(b"MZ" + b"\0" * 8)
    shutil.copy(HERE / "page.png", SAMPLES / "selfcheck_cht.png")
    shutil.copy(HERE / "text.pdf", SAMPLES / "selfcheck_text.pdf")
    shutil.copy(HERE / "sample.docx", SAMPLES / "selfcheck.docx")


if __name__ == "__main__":
    main()
